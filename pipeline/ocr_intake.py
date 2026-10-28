#!/usr/bin/env python3
"""
ocr_intake.py — Extract text from intake/ files into token-efficient JSON.

For each case under vault/cases/<case>/intake/:
  - Reads PDFs (text layer + OCR fallback), DOCX, images, .txt, .eml
  - Writes per-file sidecars: intake/_ocr/<file>.txt (full raw text)
  - Writes one index per case: intake/_ocr/index.json (compact, for LLM)

The index.json is what the model reads — small, structured, no boilerplate.
"""
import json, os, re, sys, hashlib, traceback
from pathlib import Path
from datetime import datetime

from pipeline import paths as _P
CASES_ROOT = _P.CASES
MAX_TEXT_PER_DOC = 50000    # cap per-file text in sidecar (chars)
INDEX_SNIPPET_LEN = 800     # how much text goes INTO index.json per doc
DOCKET_RE = re.compile(r'\b\d{1,2}[\-:]\d{2}[\-:]?cv[\-:]?\d{4,6}\b|\b\d{2,4}-\d{3,6}\b', re.I)
DATE_RE = re.compile(r'\b(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\.?\s+\d{1,2},?\s+\d{4}\b|\b\d{1,2}/\d{1,2}/\d{2,4}\b')
PARTY_RE = re.compile(r'\b([A-Z][a-zA-Z&.\' ]{2,40})\s+(?:v\.|vs\.|versus)\s+([A-Z][a-zA-Z&.\' ]{2,40})\b')

def extract_pdf(path):
    try:
        import pdfplumber
        text = []
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages[:50]:  # cap at 50 pages
                t = page.extract_text() or ""
                text.append(t)
        full = "\n".join(text).strip()
        if len(full) > 200:  # has real text layer
            return full[:MAX_TEXT_PER_DOC], "text"
    except Exception:
        pass
    # OCR fallback
    try:
        import pytesseract
        from pdf2image import convert_from_path
        pages = convert_from_path(str(path), dpi=200, first_page=1, last_page=10)
        ocr_text = []
        for p in pages:
            ocr_text.append(pytesseract.image_to_string(p))
        return "\n".join(ocr_text)[:MAX_TEXT_PER_DOC], "ocr"
    except Exception as e:
        return f"[OCR FAILED: {e}]", "error"

def extract_docx(path):
    try:
        from docx import Document
        doc = Document(path)
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for tbl in doc.tables:
            for row in tbl.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text)
        return "\n".join(parts)[:MAX_TEXT_PER_DOC], "text"
    except Exception as e:
        return f"[DOCX FAILED: {e}]", "error"

def extract_image(path):
    try:
        import pytesseract
        from PIL import Image
        return pytesseract.image_to_string(Image.open(path))[:MAX_TEXT_PER_DOC], "ocr"
    except Exception as e:
        return f"[IMG OCR FAILED: {e}]", "error"

def extract_text_file(path):
    try:
        return Path(path).read_text(encoding='utf-8', errors='replace')[:MAX_TEXT_PER_DOC], "text"
    except Exception as e:
        return f"[READ FAILED: {e}]", "error"

EXTRACTORS = {
    '.pdf': extract_pdf, '.PDF': extract_pdf,
    '.docx': extract_docx, '.DOCX': extract_docx, '.doc': extract_docx,
    '.png': extract_image, '.jpg': extract_image, '.jpeg': extract_image, '.tiff': extract_image, '.tif': extract_image,
    '.txt': extract_text_file, '.eml': extract_text_file, '.md': extract_text_file,
}

def file_signature(path):
    """Hash file content for change detection."""
    h = hashlib.sha256()
    with open(path, 'rb') as f:
        while chunk := f.read(65536):
            h.update(chunk)
    return h.hexdigest()[:16]

def extract_facts(text):
    """Pull out key litigation facts so the LLM doesn't have to."""
    return {
        "dockets": sorted(set(DOCKET_RE.findall(text)))[:5],
        "dates": sorted(set(DATE_RE.findall(text)))[:10],
        "parties": list({f"{m.group(1).strip()} v. {m.group(2).strip()}" for m in PARTY_RE.finditer(text)})[:3],
    }

def process_case(case_dir):
    intake_dir = case_dir / "intake"
    if not intake_dir.exists():
        return None
    
    ocr_dir = intake_dir / "_ocr"
    ocr_dir.mkdir(exist_ok=True)
    index_path = ocr_dir / "index.json"
    
    # Load existing index for incremental update
    existing = {}
    if index_path.exists():
        try:
            existing = {d["file"]: d for d in json.loads(index_path.read_text()).get("documents", [])}
        except Exception:
            existing = {}
    
    documents = []
    files = sorted([p for p in intake_dir.iterdir() if p.is_file() and p.suffix.lower() in EXTRACTORS])
    
    for fpath in files:
        rel = fpath.name
        size = fpath.stat().st_size
        sig = file_signature(fpath)
        sidecar = ocr_dir / f"{rel}.txt"
        
        # Skip if unchanged
        prev = existing.get(rel)
        if prev and prev.get("sig") == sig and sidecar.exists():
            documents.append(prev)
            continue
        
        # Extract
        ext = fpath.suffix.lower()
        extractor = EXTRACTORS[ext]
        try:
            text, source = extractor(fpath)
        except Exception as e:
            text, source = f"[EXCEPTION: {e}]", "error"
        
        # Save full text sidecar
        sidecar.write_text(text, encoding='utf-8', errors='replace')
        
        # Add compact entry to index
        documents.append({
            "file": rel,
            "size": size,
            "sig": sig,
            "extract_source": source,
            "chars": len(text),
            "snippet": text[:INDEX_SNIPPET_LEN].strip(),
            **extract_facts(text),
        })
    
    # Aggregate facts across whole case
    all_dockets, all_dates, all_parties = set(), set(), set()
    for d in documents:
        all_dockets.update(d.get("dockets", []))
        all_dates.update(d.get("dates", []))
        all_parties.update(d.get("parties", []))
    
    index = {
        "case": case_dir.name,
        "intake_path": str(intake_dir),
        "processed_at": datetime.now().isoformat(),
        "document_count": len(documents),
        "total_chars": sum(d.get("chars", 0) for d in documents),
        "all_dockets": sorted(all_dockets)[:10],
        "all_dates": sorted(all_dates)[:20],
        "all_parties": sorted(all_parties)[:10],
        "documents": documents,
    }
    index_path.write_text(json.dumps(index, indent=2))
    return index

def main():
    target_cases = sys.argv[1:] if len(sys.argv) > 1 else None
    case_dirs = sorted([d for d in CASES_ROOT.iterdir() if d.is_dir()])
    if target_cases:
        case_dirs = [d for d in case_dirs if d.name in target_cases]
    
    print(f"Processing {len(case_dirs)} cases...")
    for case_dir in case_dirs:
        try:
            t0 = datetime.now()
            idx = process_case(case_dir)
            dt = (datetime.now() - t0).total_seconds()
            if idx:
                size_kb = (case_dir / "intake" / "_ocr" / "index.json").stat().st_size // 1024
                print(f"  ✓ {case_dir.name}: {idx['document_count']} docs, "
                      f"{idx['total_chars']//1000}K chars total, "
                      f"index {size_kb}KB, {dt:.1f}s")
            else:
                print(f"  - {case_dir.name}: no intake/ folder")
        except Exception as e:
            print(f"  ✗ {case_dir.name}: {e}")
            traceback.print_exc()

if __name__ == "__main__":
    main()
