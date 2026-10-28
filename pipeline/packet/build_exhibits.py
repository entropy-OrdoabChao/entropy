#!/usr/bin/python3
"""build_exhibits.py — build the record exhibits (A–L in the ORDER_17 run) as PDFs beside the stub.
Promoted from $M/_build_exhibits.py (2026-10-25). Nothing here is argument: every file is a copy or a faithful
rendering of a record document. Rendering logic (asfiled_paragraphs / diff_runs / render, ledger table, txt->pdf) is
unchanged; the exhibit LIST comes from the stub's NEW EVIDENCE BLOCK rows (letter / title / [doc:<file>]) and the
case paths are arguments whose defaults reproduce today's ORDER_17 run.

Steps, driven by the manifest rows:
  1. --copy LETTER[.ext]=SRC    straight copies of record files to the row's [doc:] name (ext form -> <doc stem>.<ext>)
  2. every row whose [doc:] is *.txt -> <stem>.pdf via pipeline.m1._text_to_pdf; --merge LETTER=PDF appends an attachment
  3. --ledger-letter H          ledger CSV/MD -> readable PDF table at <doc stem>.pdf
  4. --clean-letter K / --comparison-letter L   amended-complaint markdown -> CLEAN .docx and Pa.R.C.P. 1033(a) COMPARISON
                                .docx (strike/underline diff against the as-filed OCR), then docx -> pdf (LibreOffice)
Run: cd <home> && env PYTHONPATH=$PWD /usr/bin/python3 -m pipeline.packet.build_exhibits --m <motion dir> [--outdir <dir>] ...
"""
from pathlib import Path
from core.entropy_paths import entropy_root as _eroot, package_root as _epkg, hermes_home as _ehermes
import argparse
import csv
import difflib
import json
import logging
import re
import shutil
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
if not (ROOT / "pipeline" / "m1.py").exists():
    ROOT = _eroot()
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")

from docx import Document  # noqa: E402
from docx.shared import Pt, Inches  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from pypdf import PdfReader, PdfWriter  # noqa: E402

from pipeline.packet.manifest import read_manifest, default_stub  # noqa: E402

# ---------------------------------------------------------------- today's values (ORDER_17 run) as defaults
DEFAULT_CASE = ROOT / "obsidian_sync/cases/example_v_insurer"
DEFAULT_M = DEFAULT_CASE / "drafts/_bundle13_order17/motions/vacate_m2room"
DEFAULT_COPIES = [
    "A={M}/motion_EXHIBIT_ORDER_CHALLENGED_signed_2026-10-21_entered_2026-10-22_ControlNo_26010000.pdf",
    "B={CASE}/drafts/_bundle11_packet/motions/chain_01_progressive_advanced/07_EXHIBIT_B_Prelim_Objections_as_filed.pdf",
    "C={CASE}/intake/PROGRESSIVE_COMPLAINT_asfiled_260100001.pdf",
    "J={CASE}/evidence/2024-08-19_ID_Cards_policy_928356298_gmail_1916c24737853f9f.pdf",
    "H.csv={CASE}/evidence/progressive_premium_ledger.csv",
    "H.md={CASE}/evidence/progressive_premium_ledger.md",
]
DEFAULT_MERGE = ["F={CASE}/evidence/2026-01-14_Progressive_reply_attachment_Filer_gmail_19bbd12973cdd237.pdf"]

TAG_RX = re.compile(r"\s*\[\s*(?:corpus|doc|official|record|source|chunk|chunk_id)\s*(?::[^\]\n]*)?\]", re.I)
MARK_RX = re.compile(r"\s*\[(?:AMENDED|NEW)\]")
SUP_RX = re.compile(r"[\u2070-\u2079\u00b9\u00b2\u00b3,]*[\u2070-\u2079\u00b9\u00b2\u00b3]")


def pages(p: Path) -> int:
    try:
        return len(PdfReader(str(p)).pages)
    except Exception:
        return 0


# ---------------------------------------------------------------- K / L renderer (unchanged from the case-local script)
def asfiled_paragraphs(text: str) -> dict:
    """{number: [paragraph text, ...]} from the OCR of the complaint as filed (numbers repeat across
    the complaint body, prayer, proof-of-facts and definitions, so each number maps to a list)."""
    out, cur_n, cur = {}, None, []

    def flush():
        if cur_n is not None and cur:
            out.setdefault(cur_n, []).append(" ".join(x.strip() for x in cur))
    for ln in text.splitlines():
        s = ln.strip()
        if not s or re.match(r"^Page \d+ of \d+$", s) or s.startswith("Case ID:"):
            continue
        m = re.match(r"^(\d+[A-Z]?)\.\s+(.*)", s)
        if m:
            flush(); cur_n, cur = m.group(1), [m.group(2)]
            continue
        if cur_n is not None and (s.isupper() or s.startswith("WHEREFORE") or s.startswith("COUNT ")):
            flush(); cur_n, cur = None, []
            continue
        if cur_n is not None:
            cur.append(s)
    flush()
    return out


def norm(s: str) -> str:
    s = SUP_RX.sub("", s)
    s = s.replace("\u201c", '"').replace("\u201d", '"').replace("\u2019", "'").replace("\u2018", "'")
    return re.sub(r"\s+", " ", s).strip()


def set_base(doc: Document):
    st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(12)
    for s in doc.sections:
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Inches(1)


def add_inline(par, text, bold=False, strike=False, underline=False):
    """**bold** aware run adder."""
    for i, chunk in enumerate(re.split(r"(\*\*[^*]+\*\*)", text)):
        if not chunk:
            continue
        b = bold
        if chunk.startswith("**") and chunk.endswith("**"):
            chunk, b = chunk[2:-2], True
        r = par.add_run(chunk); r.bold = b; r.font.strike = strike; r.font.underline = underline


def diff_runs(par, old: str, new: str):
    a, b = old.split(), new.split()
    sm = difflib.SequenceMatcher(None, a, b, autojunk=False)
    for op, i1, i2, j1, j2 in sm.get_opcodes():
        if op == "equal":
            add_inline(par, " ".join(b[j1:j2]) + " ")
        elif op == "delete":
            add_inline(par, " ".join(a[i1:i2]) + " ", strike=True)
        elif op == "insert":
            add_inline(par, " ".join(b[j1:j2]) + " ", underline=True)
        else:
            add_inline(par, " ".join(a[i1:i2]) + " ", strike=True)
            add_inline(par, " ".join(b[j1:j2]) + " ", underline=True)


def render(md_text: str, out: Path, title: str, comparison: bool, asfiled: dict = None) -> dict:
    doc = Document(); set_base(doc)
    stats = {"amended": 0, "new": 0, "unmatched": []}
    lines = md_text.splitlines()
    # drop YAML frontmatter
    if lines and lines[0].strip() == "---":
        j = next(i for i in range(1, len(lines)) if lines[i].strip() == "---")
        lines = lines[j + 1:]
    h = doc.add_paragraph(); h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run(title); r.bold = True
    if comparison:
        n = doc.add_paragraph(); n.alignment = WD_ALIGN_PARAGRAPH.CENTER
        n.add_run("Pa.R.C.P. 1033(a) comparison copy — ").italic = True
        n.add_run("deleted text struck through").font.strike = True
        n.add_run("; ").italic = True
        n.add_run("added text underlined").font.underline = True
        n.add_run(". Paragraph numbers follow the Complaint as filed; added paragraphs carry a letter suffix.").italic = True
    table_rows = []

    def flush_table():
        nonlocal table_rows
        if not table_rows:
            return
        rows = [r for r in table_rows if not set(r.replace("|", "").strip()) <= set("-: ")]
        cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
        ncol = max(len(c) for c in cells)
        t = doc.add_table(rows=len(cells), cols=ncol); t.style = "Table Grid"
        for i, row in enumerate(cells):
            for j in range(ncol):
                txt = row[j] if j < len(row) else ""
                is_new = bool(re.search(r"\[NEW\]", txt))
                txt = TAG_RX.sub("", MARK_RX.sub("", txt))
                p = t.cell(i, j).paragraphs[0]
                add_inline(p, txt, bold=(i == 0), underline=(comparison and is_new))
                for rr in p.runs:
                    rr.font.size = Pt(9)
        table_rows = []
    skip_title = True
    for ln in lines:
        s = ln.rstrip()
        if s.startswith("|"):
            table_rows.append(s); continue
        flush_table()
        if not s.strip() or s.strip() == "---":
            continue
        if s.startswith("# "):
            if skip_title:
                skip_title = False; continue
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(p, TAG_RX.sub("", s[2:]), bold=True); continue
        if s.startswith("## ") or s.startswith("### "):
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(p, TAG_RX.sub("", s.lstrip("# ")), bold=True); continue
        marked_amended = bool(re.search(r"\[AMENDED\]", s))
        marked_new = bool(re.search(r"\[NEW\]", s))
        clean = norm(TAG_RX.sub("", MARK_RX.sub("", s)))
        p = doc.add_paragraph()
        m = re.match(r"^(\d+[A-Z]?)\.\s+(.*)", clean)
        if comparison and marked_new:
            stats["new"] += 1
            add_inline(p, clean, underline=True); continue
        if comparison and marked_amended and m:
            stats["amended"] += 1
            num, body = m.group(1), m.group(2)
            cands = [norm(x) for x in (asfiled or {}).get(num, [])]
            if cands:
                old = max(cands, key=lambda c: difflib.SequenceMatcher(None, c.split(), body.split(), autojunk=False).ratio())
                add_inline(p, f"{num}. ")
                diff_runs(p, old, body)
            else:
                stats["unmatched"].append(num)
                add_inline(p, clean, underline=True)
            continue
        if comparison and marked_amended and not m:
            stats["amended"] += 1
            add_inline(p, clean, underline=True); continue
        add_inline(p, clean)
    flush_table()
    doc.save(str(out))
    return stats


# ---------------------------------------------------------------- H: ledger CSV -> readable PDF table (unchanged)
def build_ledger_pdf(h_pdf: Path, ledger_csv: Path, ledger_md: Path, csv_name: str, title_tail: str):
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib import colors
    styles = getSampleStyleSheet()
    small = styles["BodyText"]; small.fontSize = 7.5; small.leading = 9
    doc = SimpleDocTemplate(str(h_pdf), pagesize=landscape(letter), leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36,
                            title="Exhibit H — Premium ledger, Policy 928356298")
    story = [Paragraph(f"EXHIBIT H — Premium Ledger, Progressive Policy 928356298 ({title_tail})", styles["Heading2"]),
             Paragraph("Compiled from Defendant's own payment confirmation and scheduled-payment reminder emails; every row is pinned to the Gmail message id of the "
                       f"document it is taken from. Source CSV: {csv_name} (82 rows, one per month 2019-03 → 2025-12). "
                       "Totals computed with exact decimal arithmetic.", small), Spacer(1, 8)]
    md = ledger_md.read_text(encoding="utf-8")
    tot_rows = []
    for ln in md.splitlines():
        if ln.startswith("| ") and "Window" in ln or (tot_rows and ln.startswith("| ")):
            if set(ln.replace("|", "").strip()) <= set("-: "):
                continue
            cells = [Paragraph(c.strip().replace("**", ""), small) for c in ln.strip().strip("|").split("|")]
            tot_rows.append(cells)
        elif tot_rows and not ln.startswith("|"):
            break
    if tot_rows:
        t = Table(tot_rows, colWidths=[120, 190, 70, 170, 150], repeatRows=1)
        t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        story += [Paragraph("Totals", styles["Heading4"]), t, Spacer(1, 10)]
    rows = [[Paragraph(x, small) for x in ["Month / date", "Amount", "Evidence", "Gmail id", "Subject", "Note (quoted from the document)"]]]
    with open(ledger_csv, newline="", encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            rows.append([Paragraph(r.get("date_scheduled", ""), small), Paragraph(r.get("amount", ""), small),
                         Paragraph(r.get("evidence_type", ""), small), Paragraph(r.get("gmail_id", ""), small),
                         Paragraph(r.get("subject", ""), small), Paragraph((r.get("note", "") or "")[:420], small)])
    t = Table(rows, colWidths=[62, 44, 62, 92, 120, 330], repeatRows=1)
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.3, colors.grey), ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story += [Paragraph("Monthly rows", styles["Heading4"]), t]
    doc.build(story)


# ---------------------------------------------------------------- driver
def _kv(items, subs) -> dict:
    out = {}
    for it in items or []:
        k, _, v = it.partition("=")
        out[k.strip()] = Path(v.strip().format(**subs))
    return out


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="record exhibits (copies, txt->pdf, ledger, 1033 CLEAN/COMPARISON) from the stub's exhibit rows")
    ap.add_argument("--root", default=str(ROOT))
    ap.add_argument("--case", default=str(DEFAULT_CASE), help="case dir (obsidian_sync/cases/<slug>)")
    ap.add_argument("--m", default=str(DEFAULT_M), help="motion dir (stub + exhibit source files)")
    ap.add_argument("--stub", default=None, help="stub/manifest with the NEW EVIDENCE BLOCK (default <M>/EXHIBITS.md else <M>/motion_to_vacate.md)")
    ap.add_argument("--outdir", default=None, help="where exhibit files are written (default <M>)")
    ap.add_argument("--copy", action="append", default=None, help="LETTER[.ext]=SRC straight copy ({ROOT},{CASE},{M} expand); repeatable; default = today's 6")
    ap.add_argument("--merge", action="append", default=None, help="LETTER=PDF appended after that letter's txt->pdf; default F=<reply attachment>")
    ap.add_argument("--ledger-letter", default="H")
    ap.add_argument("--ledger-csv", default="{CASE}/evidence/progressive_premium_ledger.csv")
    ap.add_argument("--ledger-md", default="{CASE}/evidence/progressive_premium_ledger.md")
    ap.add_argument("--ledger-caption", default="Filer v. Example Insurer Ins. Co., Phila. CCP No. 260100001")
    ap.add_argument("--amend-md", default="{M}/EXHIBIT_PROPOSED_AMENDED_COMPLAINT.md")
    ap.add_argument("--asfiled-ocr", default="{CASE}/intake/_ocr/PROGRESSIVE_COMPLAINT_asfiled_260100001.pdf.txt")
    ap.add_argument("--clean-letter", default="K")
    ap.add_argument("--comparison-letter", default="L")
    ap.add_argument("--amend-title", default="PROPOSED FIRST AMENDED COMPLAINT")
    a = ap.parse_args(argv)

    CASE, M = Path(a.case), Path(a.m)
    subs = {"ROOT": a.root, "CASE": str(CASE), "M": str(M)}
    out = Path(a.outdir) if a.outdir else M
    out.mkdir(parents=True, exist_ok=True)
    stub = Path(a.stub) if a.stub else default_stub(M)
    rows = read_manifest(stub, bases=[M, CASE])
    by_letter = {r["letter"]: r for r in rows}
    print(f"manifest {stub.name}: {len(rows)} exhibit rows -> {out}")

    def doc_name(letter, ext=None) -> str:
        r = by_letter.get(letter)
        base = Path(r["doc"]).name if r else f"motion_EXHIBIT_{letter}.bin"
        return (Path(base).stem + "." + ext) if ext else base

    # 1. straight copies
    copies = _kv(a.copy if a.copy is not None else DEFAULT_COPIES, subs)
    for k, src in copies.items():
        letter, _, ext = k.partition(".")
        if letter not in by_letter:
            print(f"[{k}] skipped: no manifest row {letter}")
            continue
        if not src.exists():
            # DEFAULT_COPIES are one run's paths; the manifest row's own [doc:] path (resolved by
            # read_manifest against the case/record dirs) is the authority for any other bundle.
            r = by_letter[letter]; alt = Path(r["path"]) if r.get("exists") else None
            for cand in ([alt] if alt else []) + [CASE / r["doc"], M / Path(r["doc"]).name]:
                if cand and cand.exists():
                    print(f"[{k}] default source missing -> manifest path {cand}"); src = cand; break
        assert src.exists(), src
        dst = out / doc_name(letter, ext or None)
        shutil.copy2(src, dst)
        print(f"[{k}] copied -> {dst.name}")

    # 2. text exhibits -> PDF (same renderer m1 uses)
    from pipeline.m1 import _text_to_pdf
    merges = _kv(a.merge if a.merge is not None else DEFAULT_MERGE, subs)
    for r in rows:
        if not r["doc"].lower().endswith(".txt"):
            continue
        p = out / Path(r["doc"]).name
        if not p.exists():
            src = r["path"] if r["exists"] else (M / Path(r["doc"]).name)
            assert src.exists(), src
            shutil.copy2(src, p)
        pdf = _text_to_pdf(p)          # -> <name>.txt.exhibit.pdf
        final = out / (p.stem + ".pdf")
        att = merges.get(r["letter"])
        if att:
            # the reply + its attachment (Final Bill) as ONE exhibit, text first
            w = PdfWriter()
            for part in (pdf, att):
                for pg in PdfReader(str(part)).pages:
                    w.add_page(pg)
            with open(final, "wb") as fh:
                w.write(fh)
        else:
            shutil.copy2(pdf, final)
        print(f"[{r['letter']}] {final.name} pages={pages(final)}")

    # 3. ledger table
    if a.ledger_letter and a.ledger_letter in by_letter:
        h_pdf = out / doc_name(a.ledger_letter, "pdf")
        build_ledger_pdf(h_pdf, Path(a.ledger_csv.format(**subs)), Path(a.ledger_md.format(**subs)),
                         doc_name(a.ledger_letter, "csv"), a.ledger_caption)
        print(f"[{a.ledger_letter}] {h_pdf.name} pages={pages(h_pdf)}")

    # 4. K / L: amended complaint, clean + comparison -> 5. docx -> pdf (LibreOffice headless)
    amend = Path(a.amend_md.format(**subs))
    stats = {}
    if amend.exists() and (a.clean_letter in by_letter or a.comparison_letter in by_letter):
        from generation.finalize_brief import convert_to_pdf
        amend_md = amend.read_text(encoding="utf-8")
        todo = []
        if a.clean_letter in by_letter:
            k_docx = out / doc_name(a.clean_letter, "docx")
            stats[a.clean_letter] = render(amend_md, k_docx, f"EXHIBIT {a.clean_letter} — {a.amend_title} (CLEAN COPY)", comparison=False)
            print(f"[{a.clean_letter}] {k_docx.name}")
            todo.append(k_docx)
        if a.comparison_letter in by_letter:
            l_docx = out / doc_name(a.comparison_letter, "docx")
            asf = asfiled_paragraphs(Path(a.asfiled_ocr.format(**subs)).read_text(encoding="utf-8", errors="replace"))
            sl = render(amend_md, l_docx, f"EXHIBIT {a.comparison_letter} — {a.amend_title} (COMPARISON COPY)", comparison=True, asfiled=asf)
            stats[a.comparison_letter] = sl
            print(f"[{a.comparison_letter}] {l_docx.name} amended={sl['amended']} new={sl['new']} unmatched={sl['unmatched']}")
            todo.append(l_docx)
        for d in todo:
            pdf = convert_to_pdf(d)
            print(f"[pdf] {d.name} -> {pdf.name if pdf else 'FAILED'} pages={pages(pdf) if pdf else 0}")
        json.dump(stats, open(out / "_exhibit_KL_render_stats.json", "w"), indent=1)
    print("DONE")
    return 0


if __name__ == "__main__":
    sys.exit(main())
