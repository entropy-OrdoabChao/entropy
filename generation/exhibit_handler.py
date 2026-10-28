"""
generation/exhibit_handler.py
-----------------------------
Smart exhibit handler with three modes:

  REFERENCE_DOCKET   - exhibit is already on the docket (cite "Dkt N" inline, no cover)
  COVER_ONLY         - new document, generate labeled cover sheet, no source attached
  COVER_PLUS_SOURCE  - new document, generate cover sheet + append source PDF

Determines mode per exhibit by reading case_ground_truth.yaml.exhibits[] entries:

  exhibits:
    - letter: A
      title: "Proof of Service on Experian"
      mode: REFERENCE_DOCKET
      docket_ref: 54
      filed_date: "2025-07-22"
    - letter: B
      title: "Local Rule 7.1(g) Text"
      mode: COVER_ONLY
      date: "2026-05-22"
    - letter: C
      title: "Order from Related Case"
      mode: COVER_PLUS_SOURCE
      source_path: "C:/Users/malic/Documents/intake/related_order.pdf"
      date: "2026-04-15"

For REFERENCE_DOCKET exhibits, the motion body should reference them as:
  "See Dkt 54 (Exhibit A, incorporated by reference)."

For COVER_ONLY and COVER_PLUS_SOURCE, the renderer calls generate_exhibit_cover.
"""
from __future__ import annotations
from pathlib import Path
from typing import Optional
import yaml
from datetime import datetime


def classify_exhibits(case_ground_truth: dict) -> dict:
    """Split case_ground_truth.exhibits into three buckets by mode.

    Returns:
        {
            "reference_docket": [{letter, title, docket_ref, filed_date}, ...],
            "cover_only":       [{letter, title, date}, ...],
            "cover_plus_source":[{letter, title, source_path, date}, ...],
            "all":              [...full list...],
        }
    """
    exhibits = case_ground_truth.get("exhibits", []) or []
    result = {
        "reference_docket": [],
        "cover_only": [],
        "cover_plus_source": [],
        "all": exhibits,
    }
    for ex in exhibits:
        if not isinstance(ex, dict):
            continue
        mode = (ex.get("mode") or "").upper()
        if mode == "REFERENCE_DOCKET":
            result["reference_docket"].append(ex)
        elif mode == "COVER_ONLY":
            result["cover_only"].append(ex)
        elif mode == "COVER_PLUS_SOURCE":
            result["cover_plus_source"].append(ex)
        else:
            # Default behavior: if docket_ref present, treat as REFERENCE_DOCKET
            if ex.get("docket_ref"):
                ex["mode"] = "REFERENCE_DOCKET"
                result["reference_docket"].append(ex)
            elif ex.get("source_path"):
                ex["mode"] = "COVER_PLUS_SOURCE"
                result["cover_plus_source"].append(ex)
            else:
                ex["mode"] = "COVER_ONLY"
                result["cover_only"].append(ex)
    return result


def render_exhibit_list_block(classified: dict) -> str:
    """Render a markdown 'EXHIBIT LIST' section to inject into the motion body.

    Returns a markdown block that goes after the prayer for relief (or wherever
    the renderer wants exhibit references in the body).
    """
    if not classified.get("all"):
        return ""
    lines = []
    lines.append("")
    lines.append("## EXHIBIT LIST")
    lines.append("")
    for ex in classified["all"]:
        letter = ex.get("letter", "?")
        title = ex.get("title", "")
        mode = (ex.get("mode") or "").upper()
        if mode == "REFERENCE_DOCKET":
            docket = ex.get("docket_ref", "")
            filed = ex.get("filed_date", "")
            lines.append(
                f"**Exhibit {letter}** — {title}. *Filed of record at Dkt {docket}"
                + (f" ({filed})" if filed else "")
                + "; incorporated by reference. No paper copy attached.*"
            )
        elif mode == "COVER_ONLY":
            date = ex.get("date", "")
            lines.append(
                f"**Exhibit {letter}** — {title}"
                + (f" (dated {date})" if date else "")
                + ". *Cover sheet attached.*"
            )
        elif mode == "COVER_PLUS_SOURCE":
            date = ex.get("date", "")
            lines.append(
                f"**Exhibit {letter}** — {title}"
                + (f" (dated {date})" if date else "")
                + ". *Source document attached behind cover.*"
            )
        lines.append("")
    return "\n".join(lines)


def generate_covers_for_new_exhibits(
    classified: dict,
    case_id: str,
    court_id: str,
    out_dir: Path,
) -> list[dict]:
    """For COVER_ONLY and COVER_PLUS_SOURCE exhibits, generate cover sheets.

    Returns list of {letter, title, cover_path, source_path or None, mode}.
    Calls generation.md_to_docx COURT_RULES for court-aware styling, then writes
    one-page cover .docx per exhibit.
    """
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    try:
        from generation.md_to_docx import COURT_RULES
    except Exception:
        COURT_RULES = {"EDPA": {"font": "Times New Roman", "size_pt": 12,
                                "margin_in": 1.0,
                                "court_full": "UNITED STATES DISTRICT COURT\nFOR THE EASTERN DISTRICT OF PENNSYLVANIA"}}

    try:   # aliases + forum-profile courts (F2) before the EDPA styling fallback
        from generation.md_to_docx import _court_rules_or_die
        rules = _court_rules_or_die(court_id)
    except Exception:
        rules = COURT_RULES.get(court_id, COURT_RULES.get("EDPA"))
    out_dir.mkdir(parents=True, exist_ok=True)
    results = []
    targets = classified.get("cover_only", []) + classified.get("cover_plus_source", [])
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")

    for ex in targets:
        letter = (ex.get("letter") or "X").upper()
        title = ex.get("title", "")
        date = ex.get("date", "")
        source_path = ex.get("source_path")

        doc = Document()
        for s in doc.sections:
            s.page_width = Inches(8.5); s.page_height = Inches(11)
            s.top_margin = Inches(rules["margin_in"])
            s.bottom_margin = Inches(rules["margin_in"])
            s.left_margin = Inches(rules["margin_in"])
            s.right_margin = Inches(rules["margin_in"])
        style = doc.styles["Normal"]
        style.font.name = rules["font"]
        style.font.size = Pt(rules["size_pt"])

        # Court header
        for ln in rules["court_full"].split("\n"):
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(ln); r.bold = True; r.font.size = Pt(rules["size_pt"])
        # Big spacer
        for _ in range(8):
            doc.add_paragraph()
        # EXHIBIT X
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"EXHIBIT {letter}")
        r.bold = True; r.font.size = Pt(48)
        doc.add_paragraph()
        # Title
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(title); r.bold = True; r.font.size = Pt(20)
        # Date
        if date:
            doc.add_paragraph()
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(f"Dated: {date}"); r.font.size = Pt(rules["size_pt"] + 2)

        out_path = out_dir / f"exhibit_{letter}_cover_{ts}.docx"
        doc.save(str(out_path))
        results.append({
            "letter": letter,
            "title": title,
            "cover_path": str(out_path),
            "source_path": source_path,
            "mode": ex.get("mode", "COVER_ONLY"),
        })
    return results


def collect_inline_references(classified: dict) -> dict:
    """Build a dict mapping exhibit letter → inline citation string.

    For use in the motion body when referring to an exhibit:
      ref_for_dkt54 = refs["A"]  # → "Dkt 54 (Exhibit A, incorporated by reference)"
    """
    refs = {}
    for ex in classified["all"]:
        letter = ex.get("letter", "?")
        mode = (ex.get("mode") or "").upper()
        if mode == "REFERENCE_DOCKET":
            dkt = ex.get("docket_ref", "")
            refs[letter] = f"Dkt {dkt} (Exhibit {letter}, incorporated by reference)"
        else:
            refs[letter] = f"Exhibit {letter}"
    return refs
